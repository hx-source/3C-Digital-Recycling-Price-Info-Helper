from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


CELL_ADDRESS_RE = re.compile(r"^[A-Za-z]{1,3}[1-9]\d*$")


def _addresses_from_source(value: str | None) -> list[str]:
    """Get the text and optional numeric cells from importer addresses such as B8:C8."""
    if not value:
        return []
    return [
        part.strip().upper()
        for part in value.split(":")
        if CELL_ADDRESS_RE.fullmatch(part.strip())
    ]


def _merged_range_for(sheet, coordinate: str):
    for cell_range in sheet.merged_cells.ranges:
        if coordinate in cell_range:
            return cell_range
    return None


def build_excel_source_context(
    path: Path,
    sheet_name: str | None,
    cell_address: str | None,
    *,
    row_padding: int = 2,
    column_padding: int = 2,
) -> dict:
    """Build a compact, source-faithful grid around an imported Excel cell."""
    addresses = _addresses_from_source(cell_address)
    if not sheet_name or not addresses:
        raise ValueError("该记录未保存可定位的工作表或单元格地址")

    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError("原始表格中未找到对应工作表")
        sheet = workbook[sheet_name]
        locations = [
            (
                int(re.search(r"\d+$", address).group()),
                column_index_from_string(re.match(r"[A-Z]+", address).group()),
            )
            for address in addresses
        ]
        min_row = max(1, min(row for row, _ in locations) - row_padding)
        max_row = min(sheet.max_row, max(row for row, _ in locations) + row_padding)
        min_column = max(1, min(column for _, column in locations) - column_padding)
        max_column = min(sheet.max_column, max(column for _, column in locations) + column_padding)

        source_cell = addresses[0]
        price_cell = addresses[1] if len(addresses) > 1 else None
        columns = [get_column_letter(column) for column in range(min_column, max_column + 1)]
        rows = []
        for row in range(min_row, max_row + 1):
            cells = []
            for column in range(min_column, max_column + 1):
                coordinate = f"{get_column_letter(column)}{row}"
                cell = sheet[coordinate]
                merged_range = _merged_range_for(sheet, coordinate)
                value = cell.value
                if value is None and merged_range is not None:
                    value = sheet.cell(merged_range.min_row, merged_range.min_col).value
                cells.append(
                    {
                        "coordinate": coordinate,
                        "value": None if value is None else str(value),
                        "merged_range": str(merged_range) if merged_range is not None else None,
                        "is_source": coordinate == source_cell,
                        "is_price": coordinate == price_cell,
                    }
                )
            rows.append({"row": row, "cells": cells})
        return {
            "sheet_name": sheet.title,
            "source_cell": source_cell,
            "price_cell": price_cell,
            "columns": columns,
            "rows": rows,
        }
    finally:
        workbook.close()
