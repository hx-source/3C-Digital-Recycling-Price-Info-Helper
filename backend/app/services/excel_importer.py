from __future__ import annotations

import re
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from app.services.parser import ParsedCandidate, parse_text_line


EXCLUDED_SHEETS = {"思物业务一览表", "酒水"}
DATE_RE = re.compile(r"(?:(?P<year>20\d{2})[年./-])?(?P<month>\d{1,2})[月./-](?P<day>\d{1,2})")


def _sheet_date(sheet, fallback: date) -> date:
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 4), values_only=True):
        for value in row:
            if not isinstance(value, str):
                continue
            match = DATE_RE.search(value)
            if not match:
                continue
            year = int(match.group("year") or fallback.year)
            try:
                return date(year, int(match.group("month")), int(match.group("day")))
            except ValueError:
                continue
    return fallback


def _is_price_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and 10 <= float(value) <= 100000


def _best_left_text(row: tuple[Cell, ...], price_index: int) -> tuple[Cell | None, list[Cell]]:
    consumed: list[Cell] = []
    immediate = row[price_index - 1] if price_index >= 1 else None
    previous = row[price_index - 2] if price_index >= 2 else None
    if immediate and isinstance(immediate.value, str) and immediate.value.strip():
        consumed.append(immediate)
        short_code = len(immediate.value.strip()) <= 8 and not re.search(r"[\u4e00-\u9fff]{3,}", immediate.value)
        if short_code and previous and isinstance(previous.value, str) and previous.value.strip():
            consumed.append(previous)
            return previous, consumed
        return immediate, consumed
    return None, consumed


class ExcelQuoteImporter:
    def parse(self, path: Path, fallback_date: date) -> list[ParsedCandidate]:
        workbook = load_workbook(path, data_only=True, read_only=False)
        output: list[ParsedCandidate] = []

        for sheet in workbook.worksheets:
            if sheet.title in EXCLUDED_SHEETS:
                continue
            quote_date = _sheet_date(sheet, fallback_date)
            consumed_text_cells: set[str] = set()

            for row in sheet.iter_rows():
                for index, cell in enumerate(row):
                    if not _is_price_number(cell.value):
                        continue
                    text_cell, consumed = _best_left_text(row, index)
                    if text_cell is None:
                        continue
                    consumed_text_cells.update(item.coordinate for item in consumed)
                    output.extend(
                        parse_text_line(
                            str(text_cell.value),
                            sheet_name=sheet.title,
                            quote_date=quote_date,
                            cell_address=f"{text_cell.coordinate}:{cell.coordinate}",
                            source_line=text_cell.row,
                            explicit_price=Decimal(str(cell.value)),
                        )
                    )

            for row in sheet.iter_rows():
                for cell in row:
                    if cell.coordinate in consumed_text_cells or not isinstance(cell.value, str):
                        continue
                    for line_number, line in enumerate(cell.value.splitlines(), start=1):
                        output.extend(
                            parse_text_line(
                                line,
                                sheet_name=sheet.title,
                                quote_date=quote_date,
                                cell_address=cell.coordinate,
                                source_line=cell.row * 100 + line_number,
                            )
                        )
        workbook.close()
        return self._deduplicate(output)

    @staticmethod
    def _deduplicate(items: list[ParsedCandidate]) -> list[ParsedCandidate]:
        seen: set[tuple[object, ...]] = set()
        output: list[ParsedCandidate] = []
        for item in items:
            key = (
                item.sheet_name,
                item.cell_address,
                item.model_normalized.lower(),
                item.storage,
                item.color,
                item.variant,
                item.price_status,
                item.price,
            )
            if key in seen:
                continue
            seen.add(key)
            output.append(item)
        return output

